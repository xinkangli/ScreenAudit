from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
O=Path('outputs/NC_Framework_Revision_v4')
d=Document(O/'Main_Manuscript.docx')
for p in d.paragraphs:
 if p.text.startswith('The contribution is a reusable evaluation design'):
  r=p.add_run(' Supplementary Fig. 3 relates these evaluation questions to the cited literature.');r.font.name='Times New Roman';r.font.size=Pt(11)
d.save(O/'Main_Manuscript.docx')
si=Document(O/'Supplementary_Information.docx');ps=si.paragraphs;i=next(i for i,p in enumerate(ps) if p.text=='Supplementary Fig. 4');anchor=next(p._p for p in ps if p.text=='Supplementary Table 1')
for p in ps[i:i+3]:anchor.addprevious(p._p)
t=si.tables[-1];t.autofit=False
for row in t.rows:
 for col,cell in enumerate(row.cells):
  cell.width=Cm(4 if col==0 else 13.5);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
  for side in ['top','left','bottom','right']:
   el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
  pr.append(b)
t.columns[0].width=Cm(4);t.columns[1].width=Cm(13.5)
si.save(O/'Supplementary_Information.docx')

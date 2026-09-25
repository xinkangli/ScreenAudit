from pathlib import Path
from docx import Document
import re,json,hashlib
import sys
sys.path.insert(0,'work/figure_runtime')
import numpy as np,pandas as pd
O=Path('outputs/NC_Framework_Revision_v4');E=O/'scGPT_extension';runs=pd.read_csv(E/'runs.csv');tasks={t['task']:t for t in json.loads(Path('work/manuscript_data/outputs/replication_study_20260920/tasks.json').read_text())};idx=runs.set_index(['task','seed','policy']);mx=0
for line in (E/'traces.jsonl').read_text().splitlines():
 r=json.loads(line);t=tasks[r['task']];d=np.load(Path('work/manuscript_data/outputs/replication_study_20260920')/(r['task']+'.npz'));n=len(d['a']);B=min(80,max(6,int(.3*n)));initial=min(16,max(3,B//4));ids=r['selected'];assert len(ids)==len(set(ids))==B;assert r['initial']==np.random.default_rng(r['seed']).choice(n,initial,replace=False).tolist();assert ids[:initial]==r['initial'];row=idx.loc[(r['task'],r['seed'],r['policy'])]
 for m in ['a','b']:mx=max(mx,abs(float((d[m][ids]>t['threshold_'+m]).mean())-row[m]))
assert mx<1e-12
protocol=json.loads((E/'protocol.json').read_text());assert protocol['source_sha256']==hashlib.sha256((E/'run.py').read_bytes()).hexdigest()
main=Document(O/'Main_Manuscript.docx');si=Document(O/'Supplementary_Information.docx');text='\n'.join(p.text for p in main.paragraphs);assert len(main.inline_shapes)==6 and len(si.inline_shapes)==4
old=Document('outputs/NC_Second_Edition_Six_Figures/Main_Manuscript_Second_Edition.docx');refs=lambda d:[p.text for p in d.paragraphs if re.match(r'^\d+\. ',p.text)];assert refs(old)==refs(main) and len(refs(main))==40
assert not re.search('OpenAI|Use of AI tools|AI tools assisted',text)
assert '6cb5d451ab5c4b33eb673adbe4fddc61d2389df1b89b7651a9fe2e557572b922' in '\n'.join(p.text for p in si.paragraphs)
rep={'new_experiment_runs':360,'total_extension_runs':len(runs),'independent_endpoint_max_error':mx,'source_hash_verified':True,'main_figures':6,'supplementary_figures':4,'references_unchanged':40,'scGPT_coverage_eligible_tasks':12,'Tian_coverage_exclusion':['GARS','YARS'],'framework_word_counts':json.loads((O/'word_counts.json').read_text())}
(O/'validation.json').write_text(json.dumps(rep,indent=2));print(json.dumps(rep,indent=2))
